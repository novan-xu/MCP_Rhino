from __future__ import annotations

import json
import math
import queue
import re
import subprocess
import sys
import threading
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROUTER_EXE = Path(r"C:\Users\nxu\AppData\Local\MCP_Rhino\bin\MCP_Rhino.Router.exe")
DOCUMENT_PATH = r"V:\01 Project Folders\P00012 Met Tang\03-Design-Eng\03-BIM\05-Wireframe\MET - Wireframe.3dm"
OUTPUT_DIR = Path(r"V:\01 Project Folders\P00012 Met Tang\03-Design-Eng\03-BIM\05-Wireframe\04_exports\20260708_EXTR take-off by lot")
LOG_ROOT = r"C:\Projects\MCP_Rhino\Runtime_Log"
LOT_KEY = "CW_1.05_LOT"
LAYER_PATHS = [
    "01_MET-IEF-Curtain Wall::Curves - Steel",
    "01_MET-IEF-Curtain Wall::Curves - Extrusions",
    "01_MET-IEF-Curtain Wall::Curves - Stone",
]

LL_RE = re.compile(r"^\s*LL\s*(?:\*\s*([0-9]+(?:\.[0-9]+)?))?\s*$", re.IGNORECASE)


def get_ci(obj: Any, *names: str, default: Any = None) -> Any:
    if not isinstance(obj, dict):
        return default
    for name in names:
        if name in obj:
            return obj[name]
    lowered = {str(key).lower(): value for key, value in obj.items()}
    for name in names:
        found = lowered.get(name.lower())
        if found is not None:
            return found
    return default


def normalize_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def unwrap_tool_result(response: dict[str, Any], tool_name: str) -> Any:
    if "error" in response:
        raise RuntimeError(f"{tool_name} JSON-RPC error: {response['error']}")

    result = response.get("result")
    payload: Any = result
    if isinstance(result, dict) and "content" in result:
        texts: list[str] = []
        for item in result.get("content") or []:
            if isinstance(item, dict) and item.get("type") == "text":
                texts.append(str(item.get("text", "")))
        text = "\n".join(texts).strip()
        if text:
            try:
                payload = json.loads(text)
            except json.JSONDecodeError:
                payload = text

    if isinstance(payload, dict) and any(str(k).lower() == "success" for k in payload):
        success = bool(get_ci(payload, "success", "Success"))
        message = normalize_text(get_ci(payload, "message", "Message"))
        if not success:
            raise RuntimeError(f"{tool_name} failed: {message}")
        data = get_ci(payload, "data", "Data")
        return payload if data is None else data

    return payload


class McpClient:
    def __init__(self, router_exe: Path, document_path: str, timeout: float = 30.0) -> None:
        self.router_exe = router_exe
        self.document_path = document_path
        self.timeout = timeout
        self.proc: subprocess.Popen[str] | None = None
        self.stdout_queue: queue.Queue[str | None] = queue.Queue()
        self.stderr_lines: list[str] = []
        self._next_id = 1

    def __enter__(self) -> "McpClient":
        self.start()
        return self

    def __exit__(self, exc_type: Any, exc: Any, tb: Any) -> None:
        self.close()

    def start(self) -> None:
        if not self.router_exe.exists():
            raise FileNotFoundError(f"Router executable not found: {self.router_exe}")

        self.proc = subprocess.Popen(
            [str(self.router_exe)],
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
            errors="replace",
            bufsize=1,
        )

        def read_stdout() -> None:
            assert self.proc is not None and self.proc.stdout is not None
            for line in self.proc.stdout:
                if line.strip():
                    self.stdout_queue.put(line)
            self.stdout_queue.put(None)

        def read_stderr() -> None:
            assert self.proc is not None and self.proc.stderr is not None
            for line in self.proc.stderr:
                if line.strip():
                    self.stderr_lines.append(line.rstrip())

        threading.Thread(target=read_stdout, daemon=True).start()
        threading.Thread(target=read_stderr, daemon=True).start()

        self.request(
            "initialize",
            {
                "protocolVersion": "2025-11-25",
                "capabilities": {},
                "clientInfo": {"name": "met-1d-takeoff-by-lot", "version": "1.0.0"},
            },
            timeout=self.timeout,
        )
        self.notify("notifications/initialized", {})
        self.request("ping", {}, timeout=self.timeout)
        documents = self.call_tool("rhino_router_list_documents", {}, timeout=self.timeout)
        entries = get_ci(documents, "documents", default=[]) or []
        selected = next(
            (
                entry
                for entry in entries
                if normalize_text(get_ci(entry, "filePath", "FilePath")).casefold()
                == self.document_path.casefold()
            ),
            None,
        )
        if not isinstance(selected, dict):
            raise RuntimeError(f"Router did not discover the target document: {self.document_path}")
        session_id = normalize_text(get_ci(selected, "sessionId", "SessionId"))
        if not session_id:
            raise RuntimeError("Router document entry did not include a sessionId.")
        self.call_tool(
            "rhino_router_select_document",
            {"sessionId": session_id},
            timeout=self.timeout,
        )

    def close(self) -> None:
        if self.proc is None:
            return
        try:
            if self.proc.stdin and not self.proc.stdin.closed:
                self.proc.stdin.close()
            self.proc.wait(timeout=2)
        except Exception:
            if self.proc and self.proc.poll() is None:
                self.proc.kill()
        finally:
            self.proc = None

    def notify(self, method: str, params: dict[str, Any]) -> None:
        self._send({"jsonrpc": "2.0", "method": method, "params": params})

    def request(self, method: str, params: dict[str, Any], timeout: float = 120.0) -> dict[str, Any]:
        request_id = self._next_id
        self._next_id += 1
        self._send({"jsonrpc": "2.0", "id": request_id, "method": method, "params": params})
        return self._read_response(request_id, timeout)

    def call_tool(self, name: str, arguments: dict[str, Any], timeout: float = 180.0) -> Any:
        response = self.request(
            "tools/call",
            {"name": name, "arguments": arguments},
            timeout=timeout,
        )
        return unwrap_tool_result(response, name)

    def _send(self, message: dict[str, Any]) -> None:
        if self.proc is None or self.proc.stdin is None:
            raise RuntimeError("Router process is not running.")
        self.proc.stdin.write(json.dumps(message, separators=(",", ":")) + "\n")
        self.proc.stdin.flush()

    def _read_response(self, request_id: int, timeout: float) -> dict[str, Any]:
        deadline = datetime.now().timestamp() + timeout
        while True:
            remaining = deadline - datetime.now().timestamp()
            if remaining <= 0:
                stderr = "\n".join(self.stderr_lines[-10:])
                raise TimeoutError(f"Timed out waiting for JSON-RPC id {request_id}. Router stderr:\n{stderr}")
            try:
                line = self.stdout_queue.get(timeout=remaining)
            except queue.Empty as exc:
                stderr = "\n".join(self.stderr_lines[-10:])
                raise TimeoutError(f"Timed out waiting for JSON-RPC id {request_id}. Router stderr:\n{stderr}") from exc
            if line is None:
                stderr = "\n".join(self.stderr_lines[-10:])
                raise RuntimeError(f"Router closed stdout while waiting for id {request_id}. Router stderr:\n{stderr}")
            message = json.loads(line)
            if message.get("id") == request_id:
                return message


def parse_multiplier(value: str) -> float | None:
    match = LL_RE.match(value)
    if not match:
        return None
    multiplier_text = match.group(1)
    return float(multiplier_text) if multiplier_text else 1.0


def user_attributes(obj: dict[str, Any]) -> dict[str, str]:
    raw_attrs = get_ci(obj, "userAttributes", "UserAttributes", default=[]) or []
    if isinstance(raw_attrs, dict):
        return {str(key): normalize_text(value) for key, value in raw_attrs.items()}

    attrs: dict[str, str] = {}
    for entry in raw_attrs:
        if not isinstance(entry, dict):
            continue
        key = normalize_text(get_ci(entry, "key", "Key"))
        if key:
            attrs[key] = normalize_text(get_ci(entry, "value", "Value"))
    return attrs


def round_to_sixteenth(value: float) -> float:
    return math.floor((value * 16.0) + 0.5) / 16.0


def unit_to_inches_factor(unit_system: str) -> float:
    unit = unit_system.strip().lower()
    if unit in {"inch", "inches"}:
        return 1.0
    if unit in {"foot", "feet"}:
        return 12.0
    if unit in {"millimeter", "millimeters"}:
        return 1.0 / 25.4
    if unit in {"centimeter", "centimeters"}:
        return 1.0 / 2.54
    if unit in {"meter", "meters"}:
        return 39.37007874015748
    return 1.0


def simplify_part_name(key: str) -> str:
    return key[3:] if key.upper().startswith("1D-") else key


def build_workbook(rows: dict[tuple[str, str, float], float], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Takeoff"
    headers = ["Lot", "Part Name", "Length (in)", "Quantity"]
    ws.append(headers)

    for lot, part, length_in in sorted(rows.keys(), key=lambda item: (item[0], item[1], item[2])):
        qty = rows[(lot, part, length_in)]
        qty_value: int | float = int(qty) if abs(qty - round(qty)) < 1e-9 else qty
        ws.append([lot, part, length_in, qty_value])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 12
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=4):
        row[0].number_format = '0.####'
        row[1].number_format = '0.####'

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def verify_workbook(output_path: Path) -> dict[str, Any]:
    wb = load_workbook(output_path, data_only=True, read_only=True)
    if wb.sheetnames != ["Takeoff"]:
        raise RuntimeError(f"Workbook must contain one Takeoff sheet, got: {wb.sheetnames}")
    ws = wb["Takeoff"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    expected_headers = ["Lot", "Part Name", "Length (in)", "Quantity"]
    if headers != expected_headers:
        raise RuntimeError(f"Unexpected headers: {headers}")

    row_count = 0
    lot_values: set[str] = set()
    part_values: set[str] = set()
    total_quantity = 0.0
    prefixed_parts: list[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        row_count += 1
        lot = normalize_text(row[0])
        part = normalize_text(row[1])
        qty = float(row[3] or 0)
        lot_values.add(lot)
        part_values.add(part)
        total_quantity += qty
        if part.upper().startswith("1D-"):
            prefixed_parts.append(part)

    if prefixed_parts:
        raise RuntimeError(f"Part names were not simplified: {prefixed_parts[:5]}")

    return {
        "sheetNames": wb.sheetnames,
        "rowCount": row_count,
        "lotCount": len(lot_values),
        "partCount": len(part_values),
        "totalQuantity": total_quantity,
        "fileSizeBytes": output_path.stat().st_size,
    }


def main() -> int:
    if not OUTPUT_DIR.exists():
        raise FileNotFoundError(f"Output folder does not exist: {OUTPUT_DIR}")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = OUTPUT_DIR / f"MET_EXTR_1D_takeoff_by_lot_{timestamp}.xlsx"

    scope = {
        "confirmedLayerFullPaths": LAYER_PATHS,
        "objectTypes": ["Curve"],
        "matchMode": "All",
    }

    with McpClient(ROUTER_EXE, DOCUMENT_PATH, timeout=45.0) as client:
        summary = client.call_tool(
            "get_document_summary",
            {
                "filePath": DOCUMENT_PATH,
                "maxObjectSummaries": 0,
                "maxLayerSummaries": 25,
                "maxNamedViews": 0,
                "maxMaterials": 0,
            },
            timeout=90.0,
        )
        unit_system = normalize_text(get_ci(summary, "unitSystem", "UnitSystem"))
        document_name = normalize_text(get_ci(summary, "documentName", "DocumentName"))
        unit_factor = unit_to_inches_factor(unit_system)

        inspection = client.call_tool(
            "inspect_takeoff_sources",
            {"filePath": DOCUMENT_PATH, "scope": scope, "sampleValueCount": 8},
            timeout=180.0,
        )
        user_text_keys = get_ci(inspection, "userTextKeys", "UserTextKeys", default=[]) or []
        part_keys = sorted(
            {
                normalize_text(get_ci(item, "key", "Key"))
                for item in user_text_keys
                if normalize_text(get_ci(item, "key", "Key")).upper().startswith("1D")
            },
            key=str.lower,
        )
        if not part_keys:
            raise RuntimeError("No user text keys starting with 1D were found in the scoped curve layers.")

        part_conditions = [{"key": key, "comparisonMode": "Exists"} for key in part_keys]
        part_scope = {
            **scope,
            "userAttributeConditions": part_conditions,
            "userAttributeMatchMode": "Any",
        }

        filtered = client.call_tool(
            "filter_objects",
            {
                "filePath": DOCUMENT_PATH,
                "confirmedLayerFullPaths": LAYER_PATHS,
                "objectTypes": ["Curve"],
                "matchMode": "All",
                "userAttributeConditions": part_conditions,
                "userAttributeMatchMode": "Any",
            },
            timeout=300.0,
        )
        objects = get_ci(filtered, "objects", "Objects", default=[]) or []

        object_ids = [
            normalize_text(get_ci(obj, "objectId", "ObjectId"))
            for obj in objects
            if isinstance(obj, dict) and normalize_text(get_ci(obj, "objectId", "ObjectId"))
        ]
        metric_results: list[dict[str, Any]] = []
        if len(object_ids) <= 4500:
            metrics = client.call_tool(
                "get_object_metrics_by_filter",
                {
                    "filePath": DOCUMENT_PATH,
                    "confirmedLayerFullPaths": LAYER_PATHS,
                    "objectTypes": ["Curve"],
                    "matchMode": "All",
                    "userAttributeConditions": part_conditions,
                    "userAttributeMatchMode": "Any",
                },
                timeout=300.0,
            )
            metric_results.extend(get_ci(metrics, "results", "Results", default=[]) or [])
        else:
            chunk_size = 1000
            for start in range(0, len(object_ids), chunk_size):
                chunk = object_ids[start : start + chunk_size]
                metrics = client.call_tool(
                    "get_object_metrics_in_live",
                    {"filePath": DOCUMENT_PATH, "objectIds": chunk},
                    timeout=180.0,
                )
                metric_results.extend(get_ci(metrics, "results", "Results", default=[]) or [])
        lengths_by_id: dict[str, float] = {}
        metric_failures = 0
        for metric in metric_results:
            object_id = normalize_text(get_ci(metric, "objectId", "ObjectId"))
            success = bool(get_ci(metric, "success", "Success", default=False))
            length = get_ci(metric, "length", "Length")
            if object_id and success and length is not None:
                lengths_by_id[object_id.lower()] = float(length) * unit_factor
            elif object_id:
                metric_failures += 1

        grouped: dict[tuple[str, str, float], float] = defaultdict(float)
        part_object_counter: Counter[str] = Counter()
        lot_counter: Counter[str] = Counter()
        unsupported_values: Counter[tuple[str, str]] = Counter()
        missing_lot_objects = 0
        missing_metric_objects = 0

        part_key_lookup = {key.lower(): key for key in part_keys}
        for obj in objects:
            if not isinstance(obj, dict):
                continue
            object_id = normalize_text(get_ci(obj, "objectId", "ObjectId"))
            if not object_id:
                continue
            attrs = user_attributes(obj)
            attrs_lc = {key.lower(): value for key, value in attrs.items()}
            lot = normalize_text(attrs.get(LOT_KEY) or attrs_lc.get(LOT_KEY.lower()))
            if not lot:
                lot = "(missing)"
                missing_lot_objects += 1

            length_in = lengths_by_id.get(object_id.lower())
            if length_in is None:
                missing_metric_objects += 1
                continue
            rounded_length = round_to_sixteenth(length_in)

            for attr_key_lc, original_key in part_key_lookup.items():
                if attr_key_lc not in attrs_lc:
                    continue
                raw_value = attrs_lc[attr_key_lc]
                multiplier = parse_multiplier(raw_value)
                if multiplier is None:
                    unsupported_values[(original_key, raw_value)] += 1
                    continue
                part_name = simplify_part_name(original_key)
                grouped[(lot, part_name, rounded_length)] += multiplier
                part_object_counter[part_name] += 1
                lot_counter[lot] += 1

        if not grouped:
            raise RuntimeError("No grouped takeoff rows were produced.")

        build_workbook(grouped, output_path)
        verification = verify_workbook(output_path)

        client.call_tool(
            "append_activity_log",
            {
                "logRoot": LOG_ROOT,
                "task": "create MET EXTR 1D curve takeoff by CW_1.05_LOT from current live model scoped to Steel Extrusions Stone curve layers",
                "tools": [
                    "rhino_router_list_documents",
                    "rhino_router_select_document",
                    "GetDocumentSummary",
                    "InspectTakeoffSources",
                    "FilterObjects",
                    "GetObjectMetricsByFilter",
                    "GetObjectMetricsInLive",
                    "ShellCommand",
                    "AppendActivityLog",
                ],
            },
            timeout=60.0,
        )

    result = {
        "outputPath": str(output_path),
        "documentName": document_name,
        "unitSystem": unit_system,
        "scopedMatchedObjectCount": len(objects),
        "partKeyCount": len(part_keys),
        "partCount": verification["partCount"],
        "lotCount": verification["lotCount"],
        "rowCount": verification["rowCount"],
        "totalQuantity": verification["totalQuantity"],
        "missingLotObjects": missing_lot_objects,
        "missingMetricObjects": missing_metric_objects,
        "metricFailures": metric_failures,
        "unsupportedValueCount": sum(unsupported_values.values()),
        "unsupportedValueSamples": [
            {"partKey": key, "value": value, "count": count}
            for (key, value), count in unsupported_values.most_common(10)
        ],
        "topLots": lot_counter.most_common(10),
        "topParts": part_object_counter.most_common(10),
        "fileSizeBytes": verification["fileSizeBytes"],
    }
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
